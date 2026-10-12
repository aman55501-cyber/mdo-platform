"""share-analysis — entry bands, near-buy/near-sell, the Nifty 25000CE tracker, and the Share Master tabs. Canned bars, no network."""
from __future__ import annotations

import io
import os
import sys
import tempfile
from datetime import date, datetime, timedelta

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("VEGA_DB_PATH", os.path.join(tempfile.mkdtemp(), "t.db"))

import yaml  # noqa: E402

import mdo_analysis as an  # noqa: E402
from mdo_cos import IST  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _bars(prices, vol=1000.0):
    d0 = date(2026, 1, 1)
    return [{"d": (d0 + timedelta(days=i)).isoformat(), "o": p, "h": p * 1.01, "l": p * 0.99, "c": p, "v": vol} for i, p in enumerate(prices)]


def test_parse_chart_and_indicators():
    ts = [1790000000 + i * 86400 for i in range(40)]
    data = {"chart": {"result": [{"timestamp": ts, "indicators": {"quote": [{"close": [100 + i for i in range(40)], "high": [101 + i for i in range(40)],
                                                                            "low": [99 + i for i in range(40)], "open": [100 + i for i in range(40)], "volume": [1000] * 40}]}}]}}
    bars = an.parse_chart(data)
    assert len(bars) == 40 and bars[-1]["c"] == 139
    ind = an.indicators(bars)
    assert ind["last"] == 139 and ind["rsi"] == 100.0 and ind["pos52"] > 95 and abs(ind["sma20"] - 129.5) < 1e-9
    assert an.parse_chart(None) == [] and an.parse_chart({"chart": {"result": None}}) == [] and an.indicators(bars[:5]) == {}


def test_zone_parsing_and_source():
    assert an.parse_zone("Wockhardt · Mausaji (mention) · call 2026-09-23 · zone 1,965 - 2,040 · best 1,965") == (1965.0, 2040.0)
    assert an.parse_zone("no zone here") == (None, None)
    assert an.source_kind("X · Mausaji (direct) · call") == "Mausaji direct" and an.source_kind("Axis · Other (CoS)") == "Other"


def _ind(last, **kw):
    base = {"last": last, "sma20": last * 0.99, "sma50": last * 0.97, "sma200": last * 0.9, "rsi": 55.0, "atr": last * 0.02, "pos52": 50.0,
            "swing_low20": last * 0.94, "swing_high60": last * 1.2, "rvol": 1.1, "hi52": last * 1.3, "lo52": last * 0.7}
    base.update(kw)
    return base


def test_buy_side_statuses_and_entry_bands():
    z = (540.0, 550.0)
    a = an.assess(_ind(543.5), z, [], "Mausaji direct")
    assert a["status"] == "IN BUY ZONE" and a["entry_low"] == 540.0 and a["checks"] == 5 and a["sell"] == ""
    assert an.assess(_ind(560.0), z, [])["status"] == "NEAR BUY ZONE"                    # 1.8% above the top
    w = an.assess(_ind(650.0), z, [])
    assert w["status"].startswith("WAIT") and w["entry_high"] <= 650 and w["invalidation"] < w["entry_low"]
    b = an.assess(_ind(500.0), z, [])
    assert b["status"] == "BELOW ZONE" and b["entry_low"] is not None
    assert an.assess(_ind(500.0), (None, None), [])["status"].startswith("NO ZONE")


def test_sell_side_needs_a_holder_in_profit_and_never_suggests_a_loss():
    stretched = _ind(300.0, rsi=75.0)
    prof = an.assess(stretched, (None, None), [{"holder": "Ashok", "qty": 100, "avg": 250}])
    assert prof["sell"].startswith("NEAR SELL ZONE") and "Ashok" in prof["sell"]
    under = an.assess(stretched, (None, None), [{"holder": "Sudha", "qty": 100, "avg": 400}])
    assert under["sell"] == "HOLD (below cost - no loss booking)"
    unknown = an.assess(stretched, (None, None), [{"holder": "Aditi Investments", "qty": 100, "avg": 0}])
    assert unknown["sell"].startswith("HOLD (cost unknown")
    calm = an.assess(_ind(300.0), (None, None), [{"holder": "Ashok", "qty": 100, "avg": 250}])
    assert calm["sell"].startswith("HOLD (in profit")


def test_nifty_25000ce_tracker_holds_below_222_and_prices_the_option():
    day = date(2026, 10, 12)
    ns = an.nifty_status(22520.0, 14.4, day, 43.5, 29250.0, 222.0)
    assert ns["status"] == "HOLD - below buy price" and ns["days_left"] == 78 and ns["breakeven"] == 25222.0
    assert 20 < ns["model_ltp"] < 200 and 23000 < ns["need_spot"] < 26500
    assert "P&L -5,221,125" in ns["line"]
    up = an.nifty_status(26000.0, 14.4, day, None, 29250.0, 222.0)
    assert up["status"].startswith("ABOVE BUY PRICE")
    assert abs(an.bs_call(100.0, 100.0, 1.0, 0.2, 0.0) - 7.97) < 0.1


def test_run_endpoint_writes_rows_the_workbook_tabs_and_the_heartbeat(monkeypatch):
    from fastapi.testclient import TestClient
    import mdo_server
    import mdo_share_master as sm
    c = TestClient(mdo_server.app)
    c.get("/api/share-master")                                                          # schema
    rising = _bars([100 + i * 0.3 for i in range(260)])
    flat = _bars([543.5] * 260)

    def fake_hist(sym, timeout=20.0):
        return (flat if sym.startswith("HINDZINC") else (rising if sym != "^INDIAVIX" else _bars([14.4] * 40)) if sym != "^NSEI" else _bars([22520.0] * 40)), ""
    monkeypatch.setitem(mdo_server._analysis["hooks"], "fetch_history", fake_hist)

    async def fake_state(*a, **k):
        return {"portfolio": {"rows": [{"account": "Sudha", "ticker": "ADANIPOWER", "qty": 2000, "avg_price": 90.0}]},
                "levels": [{"ticker": "HINDZINC", "buy_level": 550.0, "note": "Hindustan Zinc · Mausaji (direct) · call 2026-09-22 · zone 540 - 550 · held N"}],
                "positions": [{"instrument": "NIFTY 29-Dec-26 25000 CE", "qty": 29250.0, "avg": 222.0, "ltp": 43.5, "as_of": "2026-10-09"}]}
    monkeypatch.setitem(mdo_server._share_master, "state", fake_state)
    r = c.post("/api/analysis/run", json={"now": "2026-10-12T08:40:00+05:30"}).json()
    assert r["universe"] == 2 and r["ok"] == 2 and r["in_zone"] == 1 and r["line"].startswith("share-analysis: 2 shares analysed")
    assert "Nifty 25000CE" in r["line"]
    d = c.get("/api/analysis/latest").json()
    by = {x["ticker"]: x for x in d["rows"]}
    assert by["HINDZINC"]["status"] == "IN BUY ZONE" and by["ADANIPOWER"]["holders"][0]["holder"] == "Sudha" and d["nifty"][0]["status"].startswith("HOLD")
    data = sm.build_workbook([], [], [], "n/a", analysis=d)
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    assert {"Entry Plan", "Zones", "Portfolio Check", "Nifty 25000CE"} <= set(wb.sheetnames)
    assert wb["Zones"]["A4"].value == "NEAR BUY" and wb["Zones"]["B4"].value == "HINDZINC"
    assert wb["Nifty 25000CE"]["A4"].value == "2026-10-12"
    assert "Fundamentals" in [c.value for c in wb["Entry Plan"][3]]
    wb2 = load_workbook(io.BytesIO(sm.build_workbook([], [], [], "n/a")))               # no analysis yet: tabs exist, empty, nothing invented
    assert wb2["Entry Plan"].max_row == 3


def test_fleet_cron_and_runner_exist(monkeypatch):
    import mdo_agent as ag
    bots = {b["id"]: b for b in yaml.safe_load(open(os.path.join(ROOT, "fleet.yaml"), encoding="utf-8"))["bots"]}
    assert bots["share-analysis"]["enabled"] is True and bots["share-analysis"]["provider"] == "none" and "share-analysis" in ag.CUSTOM_BOTS
    assert "mdo_agent.py share-analysis" in open(os.path.join(ROOT, "DEPLOY_HOSTINGER.md"), encoding="utf-8").read()
    beats = []
    monkeypatch.setattr(ag, "api", lambda path, method="GET", body=None, timeout=30: {"line": "share-analysis: 5 shares analysed", "status": "clean"})
    monkeypatch.setattr(ag, "heartbeat", lambda *a, **k: beats.append(a))
    assert ag.run_share_analysis({"id": "share-analysis"}, "none", "daily", now=datetime(2026, 10, 12, 8, 40, tzinfo=IST)) == 0
    assert beats[0][0] == "share-analysis" and "analysed" in beats[0][3]
