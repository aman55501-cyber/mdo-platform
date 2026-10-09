"""share-master-daily — the master share sheet: Portfolio · Mausaji Calls · Positions · Levels · Summary.

Aman, chat 2026-10-09: one workbook he can open, current every trading day with no clicks. Nothing
here touches the network: the broker bridge, Yahoo, the NSE list and Claude are canned. Follows
tests/test_levels.py: a scratch VEGA_DB_PATH is set before mdo_server is imported.
"""
from __future__ import annotations

import base64
import io
import json
import os
import sqlite3
import sys
import tempfile
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("VEGA_DB_PATH", os.path.join(tempfile.mkdtemp(), "t.db"))
os.environ.setdefault("SHARE_MASTER_DATA_DIR", tempfile.mkdtemp())

import yaml  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

import mdo_agent as ag  # noqa: E402
import mdo_share_master as sm  # noqa: E402
from mdo_cos import IST  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FRI = datetime(2026, 10, 9, 15, 45, tzinfo=IST)
BOT = {"id": "share-master-daily", "model": "claude-haiku-4-5"}

MSGS = {
    101: {"id": 101, "group_name": "Bantu Mausaji", "sender": "Bantu Mausaji", "timestamp": "2026-10-08T10:00:00+05:30",
          "text": "TCS lelo 3500 ke aas paas, target 3900, SL 3350, 3 mahine", "from_me": 0},
    102: {"id": 102, "group_name": "Bantu Mausaji", "sender": "Bantu Mausaji", "timestamp": "2026-10-09T09:10:00+05:30",
          "text": "Reliance hold karo. Zomato mat lena abhi.", "from_me": 0},
}

# synthetic HDFC Securities export: the exact header (with its repeated '% Chg.' and the extra trailing value)
HDFC_CSV = (
    "Sr. No.,Stock Name,Company Name,CMP,Change,% Change,NAV,NAV Chg,% Chg.,Portfolio Holdings,Invested Value,"
    "Average Cost Value,Unrealized Profit/Loss,% Chg,Current Value,Days Profit/Loss,% Chg.,Realized Profit/Loss,"
    "% Change,Qty,Long Term Qty,Short Term Qty,XIRR\n"
    "1,ZOMATOEQ,ETERNAL LIMITED,250.00,1.00,0.40,-,-,-,Equity,20000.0,200.0,5000.0,25.0,25000.00,-,-,0.0,-,100.0,0.0,0.0,-,-\n"
    "2,VSNLTDEQ,TATA COMMUNICATIONS LIMITED,1500.00,-10.00,-0.66,-,-,-,Equity,160000.0,1600.0,-10000.0,-6.25,150000.00,-,-,0.0,-,100.0,0.0,0.0,-,-\n"
    "3,DECGOLEQ,DECCAN GOLD MINES LTD,100.00,0,0,-,-,-,Equity,120000.0,120.0,-20000.0,-16.67,100000.00,-,-,0.0,-,1000.0,0.0,0.0,-,-\n"
    "4,XYZABCEQ,UNKNOWN WIDGETS LIMITED,10.00,0,0,-,-,-,Equity,1000.0,10.0,0.0,0.0,1000.00,-,-,0.0,-,100.0,0.0,0.0,-,-\n"
    "5,SOLDEQ,SOLD OUT LIMITED,10.00,0,0,-,-,-,Equity,-,-,-,-,-,-,-,0.0,-,0.0,0.0,0.0,-,-\n"
)
# five fake rows in NSE's EQUITY_L.csv shape (never downloaded in tests)
NSE_CSV = (
    "SYMBOL,NAME OF COMPANY, SERIES, DATE OF LISTING, PAID UP VALUE, MARKET LOT, ISIN NUMBER, FACE VALUE\n"
    "ETERNAL,Eternal Limited,EQ,23-JUL-2021,1,1,INE758T01015,1\n"
    "DECCANGOLD,Deccan Gold Mines Limited,EQ,01-JAN-2004,1,1,INE365D01011,1\n"
    "TATACOMM,Tata Communications Limited,EQ,01-JAN-2000,10,1,INE151A01013,10\n"
    "TTML,Tata Teleservices (Maharashtra) Limited,EQ,01-JAN-2000,10,1,INE517B01013,10\n"
    "RIIL,Reliance Industrial Infrastructure Limited,EQ,01-JAN-2000,10,1,INE046A01015,10\n"
)
NSE = sm.parse_nse_equity_list(NSE_CSV)

BOOK = {
    "as_of": "2026-10-09T10:15:00+00:00",
    "accounts": [
        {"creds_key": "HDFC1", "label": "HDFC1", "ok": True, "status": "ok", "reason": "", "fetched_at": "2026-10-09T10:15:00Z",
         "holdings": [{"ticker": "TCS-EQ", "quantity": 100, "average_price": 3000.0, "last_price": 3600.0},
                      {"ticker": "ZEEL", "quantity": 1000, "average_price": 200.0, "last_price": 150.0}]},
        {"creds_key": "HDFC2", "label": "HDFC2", "ok": False, "status": "degraded", "reason": "session expired — daily OTP login needed",
         "holdings": [{"ticker": "INFY", "quantity": 50, "average_price": 1400.0, "last_price": 1500.0}]},
    ],
}
POSITIONS = [
    {"account": "SA", "instrument": "BHEL 27-Oct-26 450 CE", "side": "BUY", "qty": 2625, "avg": 4.5, "ltp": 10.6, "pl": 16012.5, "product": "OVERNIGHT"},
    {"account": "SA", "instrument": "POLYCAB", "side": "SELL", "qty": -100, "avg": 8226.24, "ltp": 8270.5, "pl": -4426.0, "product": "MTF"},
    {"account": "AA", "instrument": "NIFTY 29-Dec-26 25000 CE", "side": "BUY", "qty": 29250, "avg": 222.0, "ltp": 43.5, "pl": -5221173.85, "product": "OVERNIGHT"},
    {"account": "AA", "instrument": "BANKNIFTY 13-Oct-26 FUT", "side": "SELL", "qty": -30, "avg": 56000.0, "ltp": 55900.0, "pl": 3000.0, "product": "OVERNIGHT"},
]


# ── Mausaji: prompt + parse ───────────────────────────────────────────────────
def test_parse_mausaji_calls_tolerates_garbage_and_dedups():
    raw = ("Sure, here are the calls:\n```json\n[\n"
           '{"message_id": 101, "ticker": "tcs", "action": "buy", "entry": "3,500", "target": 3900, "stop": 3350, '
           '"timeframe": "3 months", "quote": "TCS lelo 3500 ke aas paas, target 3900, SL 3350"},\n'
           '{"message_id": 101, "ticker": "TCS.NS", "action": "BUY", "quote": "repeat"},\n'          # dup → UNIQUE key
           '{"message_id": 999, "ticker": "INFY", "action": "BUY"},\n'                              # hallucinated id
           '{"message_id": 102, "ticker": "", "action": "SELL"},\n'                                 # no ticker
           '{"message_id": 102, "ticker": "RELIANCE", "action": "hold", "quote": "Reliance hold karo"},\n'
           '{"message_id": 102, "ticker": "ZOMATO", "action": "avoid", "quote": "Zomato mat lena abhi"},\n'
           '{"message_id": 102, "ticker": "SBIN", "action": "maybe"},\n'                            # not a call
           '"garbage", 42, null\n]\n```\nThat is all.')
    calls = sm.parse_mausaji_calls(raw, MSGS, "Bantu Mausaji")
    assert [(c["message_id"], c["ticker"], c["action"]) for c in calls] == [
        (101, "TCS", "BUY"), (102, "RELIANCE", "HOLD"), (102, "ZOMATO", "AVOID")]
    tcs = calls[0]
    assert (tcs["entry"], tcs["target"], tcs["stop"], tcs["timeframe"]) == (3500.0, 3900.0, 3350.0, "3 months")
    assert tcs["call_date"] == "2026-10-08" and calls[1]["call_date"] == "2026-10-09"   # from the message timestamp
    assert tcs["chat_name"] == "Bantu Mausaji" and tcs["quote"].startswith("TCS lelo")
    assert calls[1]["entry"] is None and calls[1]["target"] is None                     # never invented
    # a dict wrapper, a bare prose answer, empty, None, a truncated tail, an over-long quote
    one = sm.parse_mausaji_calls('{"calls": [{"message_id": 5, "ticker": "sbin", "action": "do not buy"}]}')
    assert one == [{**one[0], "ticker": "SBIN", "action": "AVOID"}]
    assert sm.parse_mausaji_calls("No calls today.") == [] and sm.parse_mausaji_calls("") == [] and sm.parse_mausaji_calls(None) == []
    cut = sm.parse_mausaji_calls('[{"message_id": 1, "ticker": "A", "action": "BUY"}, {"message_id": 2, "ticker": "B", "act')
    assert [c["ticker"] for c in cut] == ["A"]
    longq = sm.parse_mausaji_calls(json.dumps([{"message_id": 1, "ticker": "X", "action": "sell", "quote": "y" * 900}]))
    assert len(longq[0]["quote"]) == sm.QUOTE_MAX
    # allowed_ids as a plain set works too (no dates → today)
    assert sm.parse_mausaji_calls('[{"message_id": 7, "ticker": "X", "action": "BUY"}]', {7})[0]["message_id"] == 7
    assert sm.parse_mausaji_calls('[{"message_id": 8, "ticker": "X", "action": "BUY"}]', {7}) == []
    # the prompt carries the ids, the chat name, and the never-infer rule
    p = sm.parse_mausaji_prompt(list(MSGS.values()), "Bantu Mausaji", FRI)
    assert "[id 101]" in p and "[id 102]" in p and "Bantu Mausaji" in p and "Never infer" in p and "JSON list" in p
    assert sm.mausaji_chat() == "Bantu Mausaji"


# ── outcomes: prices only, both sides, expiry ─────────────────────────────────
def test_call_outcome_rules_both_sides_and_expiry():
    buy = {"status": "open", "action": "BUY", "target": 3900, "stop": 3350, "call_date": "2026-10-01"}
    assert sm.call_outcome(buy, 3900, FRI) == "hit_target" and sm.call_outcome(buy, 3950, FRI) == "hit_target"
    assert sm.call_outcome(buy, 3350, FRI) == "hit_stop" and sm.call_outcome(buy, 3000, FRI) == "hit_stop"
    assert sm.call_outcome(buy, 3600, FRI) is None and sm.call_outcome(buy, None, FRI) is None
    sell = {"status": "open", "action": "SELL", "target": 1300, "stop": 1450, "call_date": "2026-10-01"}
    assert sm.call_outcome(sell, 1300, FRI) == "hit_target" and sm.call_outcome(sell, 1250, FRI) == "hit_target"
    assert sm.call_outcome(sell, 1450, FRI) == "hit_stop" and sm.call_outcome(sell, 1500, FRI) == "hit_stop"
    assert sm.call_outcome(sell, 1400, FRI) is None
    hold = {"status": "open", "action": "HOLD", "target": None, "stop": None, "call_date": "2026-10-01"}
    assert sm.call_outcome(hold, 1, FRI) is None and sm.call_outcome(hold, 10 ** 6, FRI) is None
    # expiry: >60 days; a hit beats expiry; closed calls never move; no target → no hit
    assert sm.call_outcome({**buy, "call_date": "2026-08-01"}, 3600, FRI) == "expired"
    assert sm.call_outcome({**buy, "call_date": "2026-08-10"}, 3600, FRI) is None       # exactly 60 days
    assert sm.call_outcome({**buy, "call_date": "2026-08-09"}, 3600, FRI) == "expired"  # 61
    assert sm.call_outcome({**buy, "call_date": "2026-08-01"}, 3900, FRI) == "hit_target"
    assert sm.call_outcome({**buy, "status": "hit_stop"}, 3900, FRI) is None
    assert sm.call_outcome({**buy, "target": None}, 9999, FRI) is None
    out = sm.apply_outcomes([{"id": 1, **buy, "ticker": "TCS"}, {"id": 2, **sell, "ticker": "RELIANCE", "last_ltp": 1400}],
                            {"TCS": 3905}, FRI)
    assert out == [{"id": 1, "status": "hit_target", "last_ltp": 3905.0, "changed": True},
                   {"id": 2, "status": "open", "last_ltp": 1400.0, "changed": False}]   # no fresh price → keeps last_ltp


# ── portfolio rows: invested/value/pnl/weight/flags ───────────────────────────
def test_portfolio_rows_flags_weights_and_sorting():
    holdings = [{"account": "HDFC1", "ticker": "TCS-EQ", "qty": 100, "avg_price": 3000, "ltp": 3600},
                {"account": "HDFC1", "ticker": "ZEEL", "qty": 1000, "avg_price": 200, "ltp": 150},
                {"account": "ANGEL1", "ticker": "RELIANCE", "qty": 500, "avg_price": 1300, "ltp": 1400},
                {"account": "ANGEL1", "ticker": "YESBANK", "qty": 2000, "avg_price": 40, "ltp": 20},
                {"account": "ANGEL1", "ticker": "NOPRICE", "qty": 10, "avg_price": 100, "ltp": 0},
                {"account": "ANGEL1", "ticker": "", "qty": 10, "avg_price": 100}]
    rows = sm.portfolio_rows(holdings, {"tcs.ns": 3650.0})                # fetched price (any spelling) beats the broker's
    by = {r["ticker"]: r for r in rows}
    assert [r["ticker"] for r in rows] == ["RELIANCE", "YESBANK", "NOPRICE", "TCS", "ZEEL"]   # account, then value desc
    assert (by["TCS"]["ltp"], by["TCS"]["invested"], by["TCS"]["value"], by["TCS"]["pnl"], by["TCS"]["pnl_pct"]) == (3650.0, 300000.0, 365000.0, 65000.0, 21.67)
    assert (by["TCS"]["weight_pct"], by["RELIANCE"]["weight_pct"], by["ZEEL"]["weight_pct"], by["YESBANK"]["weight_pct"]) == (29.08, 55.78, 11.95, 3.19)
    assert (by["TCS"]["flag"], by["RELIANCE"]["flag"], by["ZEEL"]["flag"], by["YESBANK"]["flag"], by["NOPRICE"]["flag"]) == ("CAP", "CAP", "L15", "L30", "")
    assert by["NOPRICE"]["ltp"] is None and by["NOPRICE"]["value"] is None and by["NOPRICE"]["weight_pct"] is None   # never a guess
    assert sm.flag_for(13, -31) == "CAP+L30" and sm.flag_for(12, -15) == "L15" and sm.flag_for(None, None) == ""
    unv = sm.portfolio_rows([{"account": "A", "ticker": "DECGOL", "qty": 1, "avg_price": 1, "ltp": 0.5, "ticker_verified": False}])
    assert unv[0]["flag"] == "CAP+L30+unverified ticker"                   # one holding = 100% of the book
    joined = sm.join_levels(rows, [{"ticker": "tcs", "buy_level": 3500, "sell_level": 4200, "active": True},
                                   {"ticker": "ZEEL", "buy_level": 100, "active": False}])
    assert (by["TCS"]["buy_level"], by["TCS"]["sell_level"], by["ZEEL"]["buy_level"], by["RELIANCE"]["buy_level"]) == (3500.0, 4200.0, None, None)
    assert joined is rows


def test_holdings_from_book_never_invents_a_degraded_account():
    h, acc = sm.holdings_from_book(BOOK)
    assert [(x["account"], x["ticker"], x["qty"]) for x in h] == [("HDFC1", "TCS", 100.0), ("HDFC1", "ZEEL", 1000.0)]
    assert acc[1]["account"] == "HDFC2" and acc[1]["ok"] is False and "session expired" in acc[1]["reason"] and acc[1]["n"] == 0
    assert acc[0]["ok"] is True and acc[0]["n"] == 2
    h, acc = sm.holdings_from_book({"error": "timed out"})
    assert h == [] and acc[0]["account"] == "*" and "unreachable" in acc[0]["reason"]
    assert sm.holdings_from_book(None)[0] == []


# ── HDFC export → tickers (seed map, NSE list, unverified) ───────────────────
def test_parse_hdfc_csv_resolves_codes_and_flags_unverified():
    assert sm.hdfc_code_to_ticker("ZOMATOEQ") == ("ETERNAL", True) and sm.hdfc_code_to_ticker("ROLEXRINGSIQ") == ("ROLEXRINGS", True)
    assert sm.hdfc_code_to_ticker("DECGOLEQ") == ("DECGOL", False) and sm.hdfc_code_to_ticker("") == ("", False)
    assert NSE["count"] == 5 and NSE["by_name"]["DECCAN GOLD MINES"] == "DECCANGOLD"
    assert sm.norm_company("Tata Teleservices (Maharashtra) Limited") == "TATA TELESERVICES MAHARASHTRA"
    # seed map wins; then NSE exact; then a unique prefix (HDFC truncates); never a 4-letter stem
    assert sm.resolve_ticker("ZOMATOEQ", "ETERNAL LIMITED", NSE) == ("ETERNAL", True, "map")
    assert sm.resolve_ticker("DECGOLEQ", "DECCAN GOLD MINES LTD", NSE) == ("DECCANGOLD", True, "nse-exact")
    assert sm.resolve_ticker("RELINFRAEQ", "RELIANCE INDUSTRIAL INFRA", NSE) == ("RIIL", True, "nse-prefix")
    assert sm.resolve_ticker("XXEQ", "TATA", NSE) == ("XX", False, "unresolved")
    assert sm.resolve_ticker("DECGOLEQ", "DECCAN GOLD MINES LTD", None) == ("DECGOL", False, "unresolved")

    p = sm.parse_hdfc_csv(HDFC_CSV, "Aditi", NSE)
    assert [(h["ticker"], h["ticker_verified"], h["resolved_by"]) for h in p["holdings"]] == [
        ("ETERNAL", True, "map"), ("TATACOMM", True, "map"), ("DECCANGOLD", True, "nse-exact"), ("XYZABC", False, "unresolved")]
    assert p["unverified"] == ["XYZABC"] and p["skipped"] == [{"row": 6, "code": "SOLDEQ", "reason": "no positive qty"}]
    e = p["holdings"][0]
    assert (e["holder"], e["hdfc_code"], e["company"], e["qty"], e["avg_price"], e["cmp"], e["invested"], e["value"], e["pnl"], e["pnl_pct"]) == (
        "Aditi", "ZOMATOEQ", "ETERNAL LIMITED", 100.0, 200.0, 250.0, 20000.0, 25000.0, 5000.0, 25.0)
    assert p["resolved_by"] == {"map": 2, "nse-exact": 1, "nse-prefix": 0, "unresolved": 1}
    # without the NSE list the unmapped code stays, unverified; BOM and bytes are fine; a non-export says so
    p2 = sm.parse_hdfc_csv(("﻿" + HDFC_CSV).encode("utf-8"), "Aditi")
    assert [h["ticker"] for h in p2["holdings"]] == ["ETERNAL", "TATACOMM", "DECGOL", "XYZABC"] and p2["unverified"] == ["DECGOL", "XYZABC"]
    assert sm.parse_hdfc_csv("a,b\n1,2\n")["holdings"] == [] and "Stock Name" in sm.parse_hdfc_csv("a,b\n1,2\n")["skipped"][0]["reason"]


def test_hdfc_codes_mapped_after_2026_10_09_refresh():
    """Three of the eight codes the live refresh left unresolved are NSE-listed; the HDFC export's spelling
    of the company defeats the name normaliser, so the seed map carries them. The other five are not in
    EQUITY_L (BSE-only, an ETF, a delisting) and must stay unverified — never guessed."""
    # the exact "Company Name" strings from the 2026-10-09 exports; mapped regardless of the NSE list
    for code, company, ticker in (("AVTNATEQ", "A V T NATURAL PRODUCTS LIMITED", "AVTNPL"),
                                  ("PHICAREQ", "PCBL CHEMICALS LIMITED", "PCBL"),
                                  ("RAIALLEQ", "SARDA ENERGY & MINERRALS LTD.", "SARDAEN")):
        assert sm.hdfc_code_to_ticker(code) == (ticker, True)
        assert sm.resolve_ticker(code, company, NSE) == (ticker, True, "map")
        assert sm.resolve_ticker(code, company, None) == (ticker, True, "map")
        assert sm.resolve_by_name(company, NSE) is None               # the normaliser alone would not have caught it
    # the five that are not NSE equities keep the stripped code, unverified, with or without the list
    for code, company in (("DIATEAEQ", "DIANA TEA COMPANY LIMITED"), ("DUROFLXEQ", "VERITAS (INDIA) LIMITED"),
                          ("HDFCMFGETFEQ", "HDFC GOLD ETF"), ("JAIASSEQ", "JAIPRAKASH ASSOCIATES LIMITED"),
                          ("JSGLEASINGEQ", "COLAB PLATFORMS LIMITED")):
        stripped = code[:-2]
        assert stripped not in sm.HDFC_CODE_MAP
        assert sm.resolve_ticker(code, company, NSE) == (stripped, False, "unresolved")
    csv_text = "\n".join([
        ",".join(sm.HDFC_CSV_COLUMNS),
        "AVTNATEQ,A V T NATURAL PRODUCTS LIMITED,93.27,270479.99,90.16,9330.01,279810.00,3000.0",
        "PHICAREQ,PCBL CHEMICALS LIMITED,316.00,1962962.0,392.5924,-382962.0,1580000.00,5000.0",
        "RAIALLEQ,SARDA ENERGY & MINERRALS LTD.,483.35,1036526.19,518.2631,-69826.19,966700.00,2000.0",
        "DIATEAEQ,DIANA TEA COMPANY LIMITED,28.46,1367085.95,45.5695,-513285.95,853800.00,30000.0",
    ])
    p = sm.parse_hdfc_csv(csv_text, "Ashok", NSE)
    assert [(h["ticker"], h["ticker_verified"], h["resolved_by"]) for h in p["holdings"]] == [
        ("AVTNPL", True, "map"), ("PCBL", True, "map"), ("SARDAEN", True, "map"), ("DIATEA", False, "unresolved")]
    assert p["unverified"] == ["DIATEA"] and p["resolved_by"]["map"] == 3
    # the CoS's transcribed JSON (pf/<holder>_<date>.json): '?' marks unverified, re-resolved when the list knows the name
    doc = {"holder": "Aman", "as_of": "2026-10-09", "holdings": [
        {"code": "AFFLEEQ", "ticker": "AFFLE?", "name": "AFFLE 3I LIMITED", "qty": 500.0, "avg": 1438.75, "cmp": 1424.6, "cur": 712300.0, "pl_pct": -1.0},
        {"code": "DECGOLEQ", "ticker": "DECGOL?", "name": "DECCAN GOLD MINES LIMITED", "qty": 10, "avg": 100, "cmp": 120, "cur": 1200},
        {"code": "ZOMATOEQ", "ticker": "ETERNAL", "name": "ETERNAL LIMITED", "qty": 0, "avg": 1, "cmp": 1}]}
    j = sm.holdings_from_json(doc, nse=NSE)
    assert [(h["ticker"], h["ticker_verified"]) for h in j["holdings"]] == [("AFFLE", False), ("DECCANGOLD", True)]
    assert j["holdings"][0]["invested"] == 719375.0 and j["holdings"][0]["pnl_pct"] == -1.0 and j["holder"] == "Aman"
    assert j["skipped"][0]["code"] == "ZOMATOEQ"


def test_nse_equity_list_cache_without_network(monkeypatch, tmp_path):
    monkeypatch.setenv("SHARE_MASTER_DATA_DIR", str(tmp_path))
    calls = []
    monkeypatch.setattr(sm, "fetch_nse_equity_list", lambda timeout=20.0: calls.append(1) or NSE_CSV)
    a = sm.nse_equity_list()
    assert a["count"] == 5 and a["source"] == "download" and len(calls) == 1 and (tmp_path / sm.NSE_CACHE_FILE).exists()
    b = sm.nse_equity_list()                                                    # fresh cache → no download
    assert b["count"] == 5 and len(calls) == 1 and b["source"] == "cache"
    monkeypatch.setattr(sm, "fetch_nse_equity_list", lambda timeout=20.0: None)
    assert sm.nse_equity_list(force=True)["count"] == 5                        # download failed → stale cache
    monkeypatch.setenv("SHARE_MASTER_DATA_DIR", str(tmp_path / "empty"))
    assert sm.nse_equity_list() == {"by_name": {}, "by_isin": {}, "names": [], "count": 0, "cache_age_days": None, "source": "none"}
    assert NSE["by_isin"]["INE758T01015"] == "ETERNAL" and len(NSE["by_isin"]) == 5


# ── positions: instrument parsing, expiry, premium left ──────────────────────
def test_positions_parse_rows_and_summary():
    assert sm.parse_instrument("NIFTY 29-Dec-26 25000 CE") == {"underlying": "NIFTY", "expiry": datetime(2026, 12, 29).date(), "strike": 25000.0, "opt_type": "CE", "kind": "option"}
    assert sm.parse_instrument("BANKNIFTY 13-Oct-26 FUT")["kind"] == "future" and sm.parse_instrument("POLYCAB") == {"underlying": "POLYCAB", "expiry": None, "strike": None, "opt_type": None, "kind": "equity"}
    assert sm.parse_instrument("")["kind"] == "equity"
    rows = sm.position_rows(POSITIONS, FRI)
    by = {r["instrument"]: r for r in rows}
    assert by["BHEL 27-Oct-26 450 CE"]["days_to_expiry"] == 18 and by["BHEL 27-Oct-26 450 CE"]["premium_left"] == 27825.0
    assert by["NIFTY 29-Dec-26 25000 CE"]["premium_left"] == 1272375.0 and by["NIFTY 29-Dec-26 25000 CE"]["days_to_expiry"] == 81
    assert by["POLYCAB"]["premium_left"] is None and by["POLYCAB"]["days_to_expiry"] is None
    assert by["BANKNIFTY 13-Oct-26 FUT"]["days_to_expiry"] == 4 and by["BANKNIFTY 13-Oct-26 FUT"]["premium_left"] is None
    assert [r["instrument"] for r in rows][:2] == ["BANKNIFTY 13-Oct-26 FUT", "NIFTY 29-Dec-26 25000 CE"]   # account, nearest expiry first
    s = sm.positions_summary(rows)
    assert s["n"] == 4 and s["expiring_7d"] == 1 and s["accounts"][0] == {
        "account": "AA", "n": 2, "pl": -5218173.85, "premium_left": 1272375.0, "options": 1, "nearest_expiry": "2026-10-13", "nearest_days": 4, "as_of": None}
    good, bad = sm.validate_positions(POSITIONS + [{"account": "SA", "instrument": "X"}, "junk", {**POSITIONS[0], "side": "LONG"}])
    assert len(good) == 4 and [b["reason"][:7] for b in bad] == ["missing", "not an ", "account"]
    assert sm.fetch_positions_live() is None                                     # the HDFC hook, manual for now


# ── levels import parser ──────────────────────────────────────────────────────
def test_parse_levels_table_messy_headers():
    rows = [["Aman's list", None, None, None],                                   # a title row above the header
            ["Stock Name ", "Buy Level", " Sell  Level", "Remarks"],
            ["tcs.ns", "3,500", 4200, "IT"], ["INFY", "", "1,700", ""], ["", "", "", ""],
            ["BAD", "abc", "", ""], ["XYZ", "-5", "", ""], ["TCS", 3400, None, None]]
    out = sm.parse_levels_table(rows)
    assert out["columns"] == {"ticker": "Stock Name ", "buy": "Buy Level", "sell": " Sell  Level", "note": "Remarks"}
    assert out["levels"] == [{"ticker": "TCS", "buy_level": 3400.0, "sell_level": 4200.0, "note": "IT"},
                             {"ticker": "INFY", "buy_level": None, "sell_level": 1700.0, "note": ""}]
    assert [(s["row"], s["reason"]) for s in out["skipped"]] == [(6, "no buy or sell level"), (7, "level must be a positive number")]
    d = sm.parse_levels_table([{"Symbol": "sbin", "Entry": 800, "Target": "900", "Comment": "psu"}], default_note="x")
    assert d["levels"] == [{"ticker": "SBIN", "buy_level": 800.0, "sell_level": 900.0, "note": "psu"}] and d["columns"]["buy"] == "Entry"
    assert sm.parse_levels_table([["Company", "Buy at", "Sell above"], ["ITC", 400, 480]])["levels"][0]["ticker"] == "ITC"
    nope = sm.parse_levels_table([["Name", "Price"], ["ITC", 400]])
    assert nope["levels"] == [] and "could not find" in nope["skipped"][0]["reason"]
    assert sm.parse_levels_table([]) == {"levels": [], "skipped": [], "columns": {}}
    assert sm.rows_from_csv("a,b\n1,2\n") == [["a", "b"], ["1", "2"]]


# ── the workbook ─────────────────────────────────────────────────────────────
def _calls():
    return [{"id": 1, "call_date": "2026-10-08", "ticker": "TCS", "action": "BUY", "entry": 3500, "target": 3900, "stop": 3350,
             "timeframe": "3 months", "status": "hit_target", "ltp_at_call": 3520, "last_ltp": 3905, "last_checked": "2026-10-09T15:45",
             "quote": "TCS lelo 3500 ke aas paas", "message_id": 101, "chat_name": "Bantu Mausaji"},
            {"id": 2, "call_date": "2026-10-09", "ticker": "RELIANCE", "action": "HOLD", "entry": None, "target": None, "stop": None,
             "timeframe": "", "status": "open", "ltp_at_call": None, "last_ltp": None, "last_checked": None,
             "quote": "Reliance hold karo", "message_id": 102, "chat_name": "Bantu Mausaji"}]


def test_build_workbook_five_sheets_grouped_by_holder():
    holdings, acc = sm.holdings_from_book(BOOK)
    holdings += sm.holdings_from_csv_rows([{"holder": "Aditi", "ticker": "ETERNAL", "qty": 100, "avg_price": 200, "cmp": 250, "company": "ETERNAL LIMITED", "uploaded_at": "2026-10-09T14:00"},
                                           {"holder": "Aditi", "ticker": "XYZABC", "qty": 100, "avg_price": 10, "cmp": 9, "ticker_verified": 0, "uploaded_at": "2026-10-09T14:00"}])
    rows = sm.join_levels(sm.portfolio_rows(holdings, {"TCS": 3650}), [{"ticker": "TCS", "buy_level": 3500, "sell_level": 4200, "active": True}])
    acc[1]["snap_date"] = "2026-10-08"
    acc.append({"account": "Aditi", "ok": True, "reason": "HDFC CSV uploaded 2026-10-09 (2 holdings)", "n": 2, "source": "hdfc-csv"})
    levels = [{"ticker": "TCS", "buy_level": 3500, "sell_level": 4200, "ltp": 3650, "note": "IT", "active": True},
              {"ticker": "INFY", "buy_level": 1460, "sell_level": None, "ltp": 1455, "note": "", "active": True}]
    data = sm.build_workbook(rows, _calls(), levels, FRI, acc, positions=sm.position_rows(POSITIONS, FRI),
                             holders={"TCS": [{"holder": "HDFC1", "qty": 100}]})
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Portfolio", "Mausaji Calls", "Positions", "Levels", "Summary"]
    ws = wb["Portfolio"]
    assert ws["A1"].value == "Portfolio — as of 09 Oct 2026 15:45 IST"
    notes = [ws.cell(row=i, column=1).value for i in range(2, 6)]
    assert any("HDFC2: session expired" in str(n) and "showing last snapshot 2026-10-08" in str(n) for n in notes)
    assert any("HDFC1: live (2 holdings" in str(n) for n in notes) and any("Aditi: HDFC CSV uploaded" in str(n) for n in notes)
    hdr = next(i for i in range(1, 12) if ws.cell(row=i, column=1).value == "Holder")
    assert [ws.cell(row=hdr, column=c).value for c in (1, 2, 6, 12, 13, 14)] == ["Holder", "Ticker", "CMP", "Buy ≤", "Sell ≥", "Flag"]
    assert ws.freeze_panes == f"A{hdr + 1}" and ws.auto_filter.ref.startswith(f"A{hdr}:Q")
    assert ws.cell(row=hdr, column=17).value == "Held by (all holders)"
    col_a = [ws.cell(row=i, column=1).value for i in range(hdr + 1, hdr + 9)]
    assert col_a == ["Aditi", "Aditi", "Aditi total", "HDFC1", "HDFC1", "HDFC1 total", "ALL HOLDERS", None]   # grouped, subtotals, grand total
    tcs_row = next(i for i in range(hdr + 1, hdr + 9) if ws.cell(row=i, column=2).value == "TCS")
    assert (ws.cell(row=tcs_row, column=12).value, ws.cell(row=tcs_row, column=13).value, ws.cell(row=tcs_row, column=14).value) == (3500.0, 4200.0, "CAP")
    assert ws.cell(row=tcs_row, column=8).number_format == "#,##0.00" and ws.cell(row=tcs_row, column=10).number_format == '0.00"%"'
    assert ws.cell(row=tcs_row, column=14).fill.fgColor.rgb.endswith("FFE699")                        # CAP amber
    xyz_row = next(i for i in range(hdr + 1, hdr + 9) if ws.cell(row=i, column=2).value == "XYZABC")
    assert "unverified ticker" in ws.cell(row=xyz_row, column=14).value
    assert ws.cell(row=hdr + 3, column=8).value == 25900.0 and ws.cell(row=hdr + 7, column=8).value == 540900.0   # Aditi total, all holders
    ws = wb["Mausaji Calls"]
    assert ws["A1"].value == "Date" and ws["C1"].value == "Held by" and ws["M1"].value == "Mausaji's words" and ws.freeze_panes == "A2"
    assert ws["B2"].value == "RELIANCE" and ws["B3"].value == "TCS" and ws["I3"].value == "hit_target" and ws["M3"].value == "TCS lelo 3500 ke aas paas"
    assert ws["I3"].fill.fgColor.rgb.endswith("C6EFCE")
    assert ws["C3"].value == "HDFC1 100" and ws["C2"].value == "nobody"                         # Directive 18, from the holders map
    assert ws["K2"].value == "not checked yet"                                                   # Directive 17: never a blank price
    ws = wb["Positions"]
    assert [ws.cell(row=1, column=c).value for c in (1, 2, 11, 12)] == ["Account", "Instrument", "Days to expiry", "Premium left (qty×ltp)"]
    assert ws["A2"].value == "AA" and ws["B2"].value == "BANKNIFTY 13-Oct-26 FUT" and ws["K2"].value == 4
    assert ws["A4"].value == "AA total" and ws["H4"].value == -5218173.85 and ws["L4"].value == 1272375.0
    ws = wb["Levels"]
    assert ws["A1"].value == "Ticker" and ws["C1"].value == "Best entry" and ws["I1"].value == "Held by"
    assert ws["A2"].value == "TCS" and ws["H3"].value == "AT BUY" and ws["F2"].value == -4.11
    assert ws["I2"].value == "HDFC1 100" and ws["I3"].value == "nobody"
    ws = wb["Summary"]
    assert ws["B2"].value == "09 Oct 2026 15:45 IST" and ws["B3"].value == 4 and ws["B4"].value == "HDFC2"
    vals = [ws.cell(row=i, column=1).value for i in range(1, 60)]
    assert "Top gainers" in vals and "Top losers" in vals and "Mausaji calls" in vals and "Positions (account)" in vals and "ALL ACCOUNTS" in vals
    # empty inputs still produce a valid file with every sheet
    wb2 = load_workbook(io.BytesIO(sm.build_workbook([], [], [], "n/a")))
    assert wb2.sheetnames == ["Portfolio", "Mausaji Calls", "Positions", "Levels", "Summary"]
    assert "No broker account answered" in wb2["Portfolio"]["A2"].value and "No positions on file" in wb2["Positions"]["A2"].value


# ── endpoints through the real app ───────────────────────────────────────────
def _client():
    from fastapi.testclient import TestClient
    import mdo_server
    return TestClient(mdo_server.app)


def _wipe(c):
    for l in c.get("/api/levels").json()["levels"]:
        c.delete(f"/api/levels/{l['ticker']}")
    db = sqlite3.connect(os.environ["VEGA_DB_PATH"])
    for t in ("mausaji_calls", "portfolio_snapshot", "share_master_holdings", "share_master_positions"):
        db.execute(f"DELETE FROM {t}")
    db.execute("DELETE FROM bot_memory WHERE bot='share-master-daily'")
    db.execute("DELETE FROM whatsapp_messages WHERE lower(group_name) LIKE '%mausaji%'")
    db.commit(); db.close()


def _seed_messages():
    db = sqlite3.connect(os.environ["VEGA_DB_PATH"])
    ids = []
    for g, s, text, ts, me in (("Bantu Mausaji", "Bantu Mausaji", MSGS[101]["text"], MSGS[101]["timestamp"], 0),
                               ("Bantu Mausaji", "me", "ok mausaji", "2026-10-08T10:05:00+05:30", 1),       # from_me → never parsed
                               ("Site Group", "x", "TCS buy", "2026-10-08T10:06:00+05:30", 0),                # other chat
                               ("bantu mausaji", "Bantu Mausaji", MSGS[102]["text"], MSGS[102]["timestamp"], 0)):
        cur = db.execute("INSERT INTO whatsapp_messages (group_name, sender, text, timestamp, jid, account, chat_kind, from_me) "
                         "VALUES (?,?,?,?,?,?,?,?)", (g, s, text, ts, "91x@s.whatsapp.net", "1" if me == 0 else "2", "dm", me))
        ids.append(cur.lastrowid)
    db.commit(); db.close()
    return ids


def test_share_master_endpoints(monkeypatch, tmp_path):
    c = _client()
    _wipe(c)
    monkeypatch.setenv("SHARE_MASTER_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(sm, "fetch_nse_equity_list", lambda timeout=20.0: NSE_CSV)
    monkeypatch.setattr(sm, "fetch_book", lambda timeout=25.0: json.loads(json.dumps(BOOK)))
    prices = {"TCS": 3905.0, "ZEEL": 150.0, "ETERNAL": 260.0, "TATACOMM": None, "DECCANGOLD": 110.0, "XYZABC": None, "INFY": 1455.0}
    monkeypatch.setattr(sm, "fetch_ltps", lambda tickers: {t: prices.get(t) for t in tickers})

    # levels: import from CSV text, dry-run an xlsx, bad bodies
    r = c.post("/api/levels/import", json={"csv": "Stock,Buy Level,Sell Level\nTCS,3500,4200\nINFY,,1700\nBAD,x,\n"})
    assert r.status_code == 200 and r.json()["upserted"] == ["TCS", "INFY"] and r.json()["skipped"][0]["ticker"] == "BAD"
    wb = Workbook(); ws = wb.active; ws.append(["Symbol", "Entry", "Target"]); ws.append(["SBIN", 800, 900]); buf = io.BytesIO(); wb.save(buf)
    r = c.post("/api/levels/import", json={"xlsx_base64": base64.b64encode(buf.getvalue()).decode(), "dry_run": True})
    assert r.json()["dry_run"] is True and r.json()["levels"] == [{"ticker": "SBIN", "buy_level": 800.0, "sell_level": 900.0, "note": ""}]
    assert [l["ticker"] for l in c.get("/api/levels").json()["levels"]] == ["INFY", "TCS"]          # dry run stored nothing
    assert c.post("/api/levels/import", json={"rows": "nope"}).status_code == 400
    assert c.post("/api/levels/import", content=b"{bad").status_code == 400

    # HDFC CSV: raw text/csv with ?holder=, JSON, multipart; holder required; dry run; the holder's set is replaced
    assert c.post("/api/share-master/portfolio/import", content=HDFC_CSV, headers={"content-type": "text/csv"}).status_code == 400
    r = c.post("/api/share-master/portfolio/import?holder=Aditi", content=HDFC_CSV, headers={"content-type": "text/csv"})
    assert r.status_code == 200, r.text
    assert r.json()["count"] == 4 and r.json()["unverified"] == ["XYZABC"] and r.json()["holder"] == "Aditi"
    assert r.json()["nse_list"]["count"] == 5 and r.json()["resolved_by"]["nse-exact"] == 1
    r = c.post("/api/share-master/portfolio/import", json={"holder": " aditi ", "csv": HDFC_CSV, "dry_run": True})
    assert r.json()["dry_run"] is True and r.json()["holder"] == "aditi"
    r = c.post("/api/share-master/portfolio/import", files={"file": ("x.csv", HDFC_CSV.encode(), "text/csv")}, data={"holder": "Sudha"})
    assert r.status_code == 200 and r.json()["holder"] == "Sudha" and r.json()["stored"] == 4
    r = c.post("/api/share-master/portfolio/import", json={"holder": "Sudha", "holdings": [
        {"code": "AFFLEEQ", "ticker": "AFFLE?", "name": "AFFLE 3I LIMITED", "qty": 5, "avg": 100, "cmp": 110, "cur": 550}]})
    assert r.status_code == 200 and r.json()["count"] == 1                                           # replaced Sudha's 4 with 1
    csvs = c.get("/api/share-master/portfolio/csv").json()
    assert csvs["count"] == 5 and set(csvs["holders"]) == {"Aditi", "Sudha"} and csvs["holders"]["Aditi"]["unverified"] == 1
    assert c.post("/api/share-master/portfolio/import", json={"holder": "A", "csv": "a,b\n1,2\n"}).status_code == 400

    # positions: exact keys, replace-all per account, dry run
    r = c.post("/api/share-master/positions/import", json={"positions": POSITIONS, "as_of": "2026-10-09 14:15 IST"})
    assert r.status_code == 200 and r.json()["stored"] == 4 and r.json()["accounts"] == ["AA", "SA"]
    r = c.post("/api/share-master/positions/import", json={"positions": [POSITIONS[0]], "as_of": "2026-10-09 15:30 IST"})
    pos = c.get("/api/share-master/positions").json()
    assert pos["count"] == 3 and [p["account"] for p in pos["positions"]] == ["AA", "AA", "SA"]     # SA replaced, AA kept
    assert pos["positions"][2]["as_of"] == "2026-10-09 15:30 IST" and pos["positions"][2]["premium_left"] == 27825.0
    assert c.post("/api/share-master/positions/import", json={"positions": [{"account": "SA"}]}).status_code == 400
    assert c.post("/api/share-master/positions/import", json={"positions": []}).status_code == 400
    assert c.post("/api/share-master/positions/import", json={"positions": POSITIONS[:1], "dry_run": True}).json()["valid"] == 1

    # before any refresh: the CSV holders are named, nothing else is claimed
    st = c.get("/api/share-master").json()
    assert st["portfolio"]["rows"] == [] and st["portfolio"]["snap_date"] is None
    assert {a["account"]: a["ok"] for a in st["portfolio"]["accounts"]} == {"Aditi": True, "Sudha": True}
    assert st["summary"]["total"]["n"] == 0 and st["mausaji_chat"] == "Bantu Mausaji" and st["summary"]["positions"]["n"] == 3

    # calls insert dedups on (message_id, ticker, action); rejects junk
    r = c.post("/api/share-master/calls", json={"calls": [
        {"message_id": 101, "ticker": "TCS", "action": "BUY", "entry": 3500, "target": 3900, "stop": 3350, "quote": "TCS lelo", "call_date": "2026-10-08"},
        {"message_id": 101, "ticker": "tcs.ns", "action": "buy", "quote": "again"},
        {"message_id": 102, "ticker": "RELIANCE", "action": "SELL", "target": 1300, "stop": 1450, "quote": "bech do", "call_date": "2026-10-09"},
        {"message_id": 90, "ticker": "OLD", "action": "BUY", "call_date": "2026-07-01", "quote": "purana"},
        {"ticker": "NOID", "action": "BUY"}, "junk"]})
    assert r.status_code == 200 and (r.json()["inserted"], r.json()["duplicates"], r.json()["rejected"]) == (3, 1, 2)
    assert c.post("/api/share-master/calls", json={"calls": "x"}).status_code == 400

    # refresh: HDFC1 live, HDFC2 stale (no rows invented), Aditi + Sudha from CSV, prices, outcomes, workbook → vault
    r = c.post("/api/share-master/refresh", json={})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["stored"] == 7 and d["snap_date"] == datetime.now(IST).date().isoformat()
    acc = {a["account"]: a for a in d["portfolio"]["accounts"]}
    assert acc["HDFC1"]["ok"] is True and acc["HDFC2"]["ok"] is False and acc["Aditi"]["ok"] is True and acc["Sudha"]["ok"] is True
    assert acc["HDFC2"]["note"] == "HDFC2: session expired — daily OTP login needed — no snapshot on file"
    assert acc["Aditi"]["note"].startswith("Aditi: HDFC CSV uploaded") and "fallback" in acc["Aditi"]["note"] and "1 code(s) not in the HDFC→NSE map" in acc["Aditi"]["note"]
    rows = {(x["account"], x["ticker"]): x for x in d["portfolio"]["rows"]}
    assert ("HDFC2", "INFY") not in rows and len(rows) == 7
    assert rows[("HDFC1", "TCS")]["ltp"] == 3905.0 and rows[("HDFC1", "TCS")]["buy_level"] == 3500.0 and rows[("HDFC1", "TCS")]["flag"] == "CAP"
    assert rows[("Aditi", "TATACOMM")]["ltp"] == 1500.0 and rows[("Aditi", "XYZABC")]["flag"] == "unverified ticker"   # avg 10, CSV CMP 10 → no loss flag; Yahoo n/a → the CSV's CMP
    assert rows[("Aditi", "ETERNAL")]["company"] == "ETERNAL LIMITED"
    assert d["outcomes"] == {"hit_target": 1, "hit_stop": 0, "expired": 1, "checked": 3}
    # holdings without a price (Sudha's AFFLE came in via JSON; Yahoo stub has none) — never the levels/calls tickers
    assert d["ltp_missing"] == ["AFFLE", "TATACOMM", "XYZABC"]
    assert not set(d["ltp_missing_other"]) & set(d["ltp_missing"]) and "AFFLE" not in d["ltp_missing_other"]
    assert d["vault"]["saved"] is True and d["vault"]["path"] == "finance/Share_Master.xlsx"
    vault_file = os.path.join(sm.vault_dir(), "finance", "Share_Master.xlsx")
    assert os.path.isfile(vault_file) and not os.path.exists(vault_file + ".prev")
    calls = {x["ticker"]: x for x in c.get("/api/share-master/calls").json()["calls"]}
    assert calls["TCS"]["status"] == "hit_target" and calls["TCS"]["last_ltp"] == 3905.0 and calls["TCS"]["ltp_at_call"] == 3905.0
    assert calls["RELIANCE"]["status"] == "open" and calls["OLD"]["status"] == "expired"
    assert c.get("/api/share-master/calls?status=hit_target").json()["count"] == 1 and c.get("/api/share-master/calls?status=zzz").status_code == 400
    st = c.get("/api/share-master").json()                                     # the GET shows what the last refresh saw
    assert {a["account"]: a["ok"] for a in st["portfolio"]["accounts"]} == {"HDFC1": True, "HDFC2": False, "Aditi": True, "Sudha": True}
    assert st["summary"]["total"]["n"] == 7 and st["summary"]["stale_accounts"] == ["HDFC2"] and st["summary"]["unverified"] == 2
    assert st["summary"]["calls"] == {"open": 1, "hit_target": 1, "hit_stop": 0, "expired": 1, "total": 3}
    assert st["summary"]["levels"] == 2 and st["portfolio"]["rows"][0]["buy_level"] in (3500.0, None)

    # a second refresh keeps the previous workbook as .prev; a live account beats its CSV copy
    monkeypatch.setattr(sm, "fetch_book", lambda timeout=25.0: {"as_of": "x", "accounts": [
        {"creds_key": "HDFC3", "label": "Aditi", "ok": True, "status": "ok", "holdings": [{"ticker": "ETERNAL", "quantity": 100, "average_price": 200, "last_price": 255}]}]})
    d = c.post("/api/share-master/refresh", json={}).json()
    assert os.path.exists(vault_file + ".prev")
    acc = {a["account"]: a for a in d["portfolio"]["accounts"]}
    assert acc["Aditi"]["ok"] is True and acc["Aditi"]["note"] == "Aditi: live, 1 holdings" and acc["HDFC1"]["ok"] is False
    assert "last snapshot" in acc["HDFC1"]["note"]
    assert sum(1 for x in d["portfolio"]["rows"] if x["account"] == "Aditi") == 1               # live, not the 4 CSV rows
    # bridge down entirely → every broker account stale, CSV holders still served, nothing invented
    monkeypatch.setattr(sm, "fetch_book", lambda timeout=25.0: {"error": "connection refused"})
    d = c.post("/api/share-master/refresh", json={"build": False}).json()
    acc = {a["account"]: a for a in d["portfolio"]["accounts"]}
    assert acc["HDFC1"]["ok"] is False and "Shares CFO unreachable" in acc["HDFC1"]["note"] and acc["Aditi"]["ok"] is True
    assert "vault" not in d

    # the download is a real workbook with every sheet
    r = c.get("/api/share-master/xlsx")
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats") and "Share_Master_" in r.headers["content-disposition"]
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Portfolio", "Mausaji Calls", "Positions", "Levels", "Summary"]

    # Mausaji messages: the chat by display name (any case, both accounts), his side only, after the watermark
    ids = _seed_messages()
    m = c.get("/api/share-master/mausaji/messages").json()
    assert m["chat"] == "Bantu Mausaji" and [x["id"] for x in m["messages"]] == [ids[0], ids[3]] and m["chat_total"] == 3
    assert c.get(f"/api/share-master/mausaji/messages?since_id={ids[0]}").json()["messages"][0]["id"] == ids[3]
    monkeypatch.setenv("MAUSAJI_CHAT", "nobody")
    assert c.get("/api/share-master/mausaji/messages").json()["count"] == 0
    _wipe(c)


# ── the bot, end to end: canned Claude, canned broker, canned prices ──────────
class _Fake:
    def __init__(self, client):
        self.client, self.reports, self.memory, self.calls = client, [], {}, []

    def api(self, path, method="GET", body=None, timeout=30):
        self.calls.append((method, path))
        if path == "/api/agent/report":
            self.reports.append(body)
            return {"stored": True, "report_id": len(self.reports)}
        if path == "/api/cos/memory":
            if method == "POST":
                self.memory[(body["bot"], body["key"])] = body
                return {"stored": True}
            return {"bot_memory": [{"bot": b, "key": k, "value": v["value"], "status": "active"} for (b, k), v in self.memory.items()]}
        if path.startswith("/api/share-master"):
            r = self.client.post(path, json=body) if method == "POST" else self.client.get(path)
            assert r.status_code == 200, r.text
            return r.json()
        raise AssertionError(f"unexpected call {method} {path}")


def test_run_share_master_bot_end_to_end(monkeypatch, tmp_path):
    c = _client()
    _wipe(c)
    ids = _seed_messages()
    monkeypatch.setenv("SHARE_MASTER_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(sm, "fetch_book", lambda timeout=25.0: json.loads(json.dumps(BOOK)))
    monkeypatch.setattr(sm, "fetch_ltps", lambda tickers: {t: {"TCS": 3905.0, "ZEEL": 150.0, "RELIANCE": 1400.0}.get(t) for t in tickers})
    c.post("/api/levels", json=[{"ticker": "TCS", "buy_level": 3500, "sell_level": 4200}])
    fake = _Fake(c)
    monkeypatch.setattr(ag, "api", fake.api)
    prompts: list[str] = []

    def canned(prompt, model, bot_id, max_tokens=16000):
        prompts.append(prompt)
        assert model == "claude-haiku-4-5" and bot_id == "share-master-daily"
        return json.dumps([
            {"message_id": ids[0], "ticker": "TCS", "action": "BUY", "entry": 3500, "target": 3900, "stop": 3350, "timeframe": "3 mahine", "quote": "TCS lelo 3500 ke aas paas, target 3900, SL 3350"},
            {"message_id": ids[3], "ticker": "RELIANCE", "action": "HOLD", "quote": "Reliance hold karo"},
            {"message_id": ids[3], "ticker": "ZOMATO", "action": "AVOID", "quote": "Zomato mat lena abhi"},
            {"message_id": 999999, "ticker": "FAKE", "action": "BUY", "quote": "not in the batch"}])
    monkeypatch.setattr(ag, "ask_claude", canned)

    assert ag.run_share_master(BOT, "claude-haiku-4-5", "daily", now=FRI) == 0
    assert len(prompts) == 1 and f"[id {ids[0]}]" in prompts[0] and f"[id {ids[3]}]" in prompts[0]
    assert f"[id {ids[1]}]" not in prompts[0] and "TCS buy" not in prompts[0]                  # not his side, not that chat
    rep = fake.reports[-1]
    assert rep["heartbeat"] is True and rep["bot"] == "share-master-daily" and rep["status"] == "warning"   # HDFC2 stale → warning
    line = rep["summary"]
    assert line.startswith("share-master: 2 holdings across 2 accounts (stale: HDFC2) · 3 Mausaji calls (+3 new calls, 1 hit target, 0 stopped) · levels 1")
    assert "Mausaji chat 'Bantu Mausaji': 2 new msg(s) read, 3 call(s) found" in line and "saved finance/Share_Master.xlsx" in line
    assert fake.memory[("share-master-daily", sm.WATERMARK_KEY)]["value"] == str(ids[3])
    calls = {x["ticker"]: x for x in c.get("/api/share-master/calls").json()["calls"]}
    assert set(calls) == {"TCS", "RELIANCE", "ZOMATO"} and calls["TCS"]["status"] == "hit_target" and calls["TCS"]["quote"].startswith("TCS lelo")
    assert calls["RELIANCE"]["call_date"] == "2026-10-09"

    # second run: nothing new → no Claude call, says so; stale account still named
    assert ag.run_share_master(BOT, "claude-haiku-4-5", "daily", now=FRI) == 0
    assert len(prompts) == 1
    line = fake.reports[-1]["summary"]
    assert f"Mausaji chat 'Bantu Mausaji': no new messages since #{ids[3]}" in line and "(+0 new calls, 0 hit target" in line
    # the chat never stored at all → the heartbeat asks whether it is classified personal
    monkeypatch.setenv("MAUSAJI_CHAT", "Someone Else")
    assert ag.run_share_master(BOT, "claude-haiku-4-5", "daily", now=FRI) == 0
    assert "chat never seen in whatsapp_messages" in fake.reports[-1]["summary"]
    monkeypatch.delenv("MAUSAJI_CHAT")
    # Claude down → the watermark does not advance past the failed batch; the refresh still runs
    _wipe(c)
    ids = _seed_messages()
    fake.memory.clear()
    monkeypatch.setattr(ag, "ask_claude", lambda *a, **k: (_ for _ in ()).throw(RuntimeError("api down")))
    assert ag.run_share_master(BOT, "claude-haiku-4-5", "daily", now=FRI) == 0
    assert ("share-master-daily", sm.WATERMARK_KEY) not in fake.memory and "parse batch 1: RuntimeError" in fake.reports[-1]["summary"]
    assert "0 Mausaji calls" in fake.reports[-1]["summary"]
    # backend refresh failing → error heartbeat, exit 3
    def broken(path, method="GET", body=None, timeout=30):
        if path == "/api/share-master/refresh":
            raise RuntimeError("500 boom")
        return fake.api(path, method, body, timeout)
    monkeypatch.setattr(ag, "api", broken)
    monkeypatch.setattr(ag, "ask_claude", canned)
    assert ag.run_share_master(BOT, "claude-haiku-4-5", "daily", now=FRI) == 3
    assert fake.reports[-1]["status"] == "error" and "refresh failed" in fake.reports[-1]["summary"]
    _wipe(c)


def test_share_master_wired_in_dispatch_fleet_cron_env_and_requirements():
    assert ag.CUSTOM_BOTS["share-master-daily"] is ag.run_share_master
    fleet = yaml.safe_load(open(os.path.join(ROOT, "fleet.yaml"), encoding="utf-8"))
    bot = ag.find_bot(fleet, "share-master-daily")
    assert bot["enabled"] is True and bot["runs_on"] == "vps-cron" and bot["model"] == "claude-haiku-4-5"
    assert bot["command"] == "docker compose exec -T backend python mdo_agent.py share-master-daily"
    assert bot["cadence"] == "daily 15:45 IST + daily 08:45 IST" and bot["serves"] == ["capital-rules"] and bot["heartbeat"]
    assert {"purpose", "thinks", "works", "limits", "minimum_output", "rules"} <= set(bot["charter"])
    rules = " ".join(bot["charter"]["rules"]).lower()
    assert "never invents a holding or a call" in rules and "never suggests execution" in rules
    assert "mausaji's own words" in rules and "prices only" in rules
    doc = open(os.path.join(ROOT, "DEPLOY_HOSTINGER.md"), encoding="utf-8").read()
    block = doc[doc.index("# MDO fleet — one line per bot"):]
    block = block[:block.index("```")]
    lines = [ln for ln in block.splitlines() if "mdo_agent.py share-master-daily" in ln]
    assert [ln.split()[:5] for ln in lines] == [["15", "10", "*", "*", "1-5"], ["15", "3", "*", "*", "1-5"]]
    env = open(os.path.join(ROOT, ".env.example"), encoding="utf-8").read()
    assert "MAUSAJI_CHAT=Bantu Mausaji" in env and "SHARESCFO_URL" in env
    req = open(os.path.join(ROOT, "requirements_server.txt"), encoding="utf-8").read()
    assert "openpyxl" in req and "python-multipart" in req
    assert os.path.isfile(os.path.join(ROOT, "tools", "import_share_master.py"))
    assert sm.sharescfo_url() in ("http://sharescfo:8000", os.environ.get("CFO_API_URL", "").rstrip("/"), os.environ.get("SHARESCFO_URL", "").rstrip("/"))


if __name__ == "__main__":
    import pytest
    sys.exit(pytest.main([__file__, "-q"]))


# ── Angel One "DP Transaction Cum Holding" statement (A23) ───────────────────
# Three scripts as pdftotext -layout prints them: ETERNAL's block is cut by a page break (\f) and its
# ISIN header repeats on page 2 — the LAST closing line wins (150, not 100). TATACOMM resolves by ISIN;
# the third ISIN is not in the fixture list, so its name resolves it (TTML, nse-exact). A sold-out
# script (closing qty 0) and an unknown name are listed as skipped / unverified, never shown as held.
ANGEL_TXT = (
    "                  ANGEL ONE LIMITED - DP TRANSACTION CUM HOLDING STATEMENT\n"
    "Client: ADITI INVESTMENTS      DP ID: 12033200      Period: 01-Oct-2026 to 08-Oct-2026\n"
    "\n"
    "INE758T01015    ETERNAL-EQ\n"
    "Date          Description                      Debit        Credit       Balance      Value\n"
    "01-Oct-2026   OPENING BALANCE                                            100.00\n"
    "              CLOSING BALANCE                  0.00         0.00         100.00       26,000.00\n"
    "INE151A01013    TATA COMMUNICATIONS-EQ\n"
    "              CLOSING BALANCE                  0.00         0.00         10.00        15,000.00\n"
    "                                   Page 1 of 2\n"
    "\f"
    "INE758T01015    ETERNAL-EQ\n"
    "05-Oct-2026   CREDIT - MARKET PURCHASE                      50.00\n"
    "              CLOSING BALANCE                  0.00         50.00        150.00       39,000.00\n"
    "INE517B01013    TATA TELESERVICES (MAHARASHTRA)-EQ\n"
    "              CLOSING BALANCE                  0.00         0.00         2,000.00     1,60,000.00\n"
    "INE000X01099    SOLD OUT CO-EQ\n"
    "              CLOSING BALANCE                  500.00       0.00         0.00         0.00\n"
    "INE999Z01999    SOME UNKNOWN CO-EQ\n"
    "              CLOSING BALANCE                  0.00         0.00         5.00         1,000.00\n"
    "INE888Y01888    BROKEN LINE CO-EQ\n"
    "              CLOSING BALANCE                  12.00\n"
    "                                   Page 2 of 2\n"
)
NSE_NO_TTML_ISIN = sm.parse_nse_equity_list(NSE_CSV.replace("INE517B01013", "INE517B01999"))   # TTML only by name


def test_parse_angel_dp_text():
    p = sm.parse_angel_dp_text(ANGEL_TXT, "Aditi Investments", NSE_NO_TTML_ISIN)
    rows = {h["isin"]: h for h in p["holdings"]}
    assert [h["ticker"] for h in p["holdings"]] == ["ETERNAL", "TATACOMM", "TTML", "INE999Z01999"] and p["scripts"] == 6
    e = rows["INE758T01015"]
    assert e["qty"] == 150.0 and e["value"] == 39000.0 and e["cmp"] == 260.0 and e["closing_lines"] == 2       # page break: last wins
    assert e["resolved_by"] == "nse-isin" and e["ticker_verified"] is True and e["broker"] == "Angel One" and e["holder"] == "Aditi Investments"
    assert e["avg_price"] is None and e["invested"] is None and e["pnl"] is None and e["source"] == "angel-dp"  # DP statement: no cost
    assert rows["INE151A01013"]["qty"] == 10.0 and rows["INE151A01013"]["cmp"] == 1500.0
    t = rows["INE517B01013"]
    assert t["ticker"] == "TTML" and t["resolved_by"] == "nse-exact" and t["qty"] == 2000.0 and t["value"] == 160000.0 and t["cmp"] == 80.0
    u = rows["INE999Z01999"]
    assert u["ticker"] == "INE999Z01999" and u["ticker_verified"] is False and p["unverified"] == ["INE999Z01999"]
    assert [(x["isin"], x["reason"]) for x in p["skipped"]] == [
        ("INE000X01099", "closing qty 0.0 — not held"),
        ("INE888Y01888", "line 22: CLOSING BALANCE has 1 number(s), need debit credit qty amount")]
    assert p["resolved_by"] == {"nse-isin": 2, "nse-exact": 1, "nse-prefix": 0, "unresolved": 1}
    # with the full list TTML resolves by ISIN; bytes are fine; a non-statement parses to nothing, says so
    assert sm.parse_angel_dp_text(ANGEL_TXT.encode(), "x", NSE)["resolved_by"]["nse-isin"] == 3
    assert sm.parse_angel_dp_text("hello\nCLOSING BALANCE 1 2 3 4\n")["holdings"] == []                 # no ISIN header → no script
    assert sm.resolve_ticker_isin("INE000000000", "TATA", NSE) == ("INE000000000", False, "unresolved")  # 'TATA' alone never resolves


def _tiny_pdf(pages: list[list[str]]) -> bytes:
    """A minimal PDF (Courier, one text line per string) — enough for pdftotext and pypdf to read."""
    objs: list[bytes] = []

    def add(b: bytes) -> int:
        objs.append(b)
        return len(objs)

    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>")
    pages_id = add(b"")
    page_ids = []
    for lines in pages:
        esc = lambda l: l.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")  # noqa: E731
        content = ("BT /F1 9 Tf 11 TL 30 800 Td " + " ".join(f"({esc(l)}) Tj T*" for l in lines) + " ET").encode("latin-1")
        cid = add(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
        page_ids.append(add(b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 %d 0 R >> >> "
                            b"/Contents %d 0 R >>" % (pages_id, font, cid)))
    objs[pages_id - 1] = (b"<< /Type /Pages /Kids [" + b" ".join(b"%d 0 R" % i for i in page_ids) + b"] /Count %d >>" % len(page_ids))
    cat = add(b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id)
    out, offs = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, start=1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, cat, xref)
    return bytes(out)


def test_angel_pdf_import_endpoint(monkeypatch, tmp_path):
    c = _client()
    _wipe(c)
    monkeypatch.setenv("SHARE_MASTER_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(sm, "fetch_nse_equity_list", lambda timeout=20.0: NSE_CSV)
    pages = [blk.splitlines() for blk in ANGEL_TXT.split("\f")]
    pdf = _tiny_pdf(pages)
    # pdf_to_text: poppler's pdftotext when present, else pypdf — both keep the layout the parser needs
    txt = sm.pdf_to_text(pdf)
    assert "INE758T01015" in txt and "CLOSING BALANCE" in txt
    monkeypatch.setattr(sm.shutil, "which", lambda name: None)
    assert "INE517B01013" in sm.pdf_to_text(pdf)                                                  # pypdf path
    with __import__("pytest").raises(ValueError):
        sm.pdf_to_text(b"%PDF-1.4 not really")

    # the PDF upload, broker=angel (multipart), holder required, dry run first
    r = c.post("/api/share-master/portfolio/import", files={"file": ("Aditi_DP.pdf", pdf, "application/pdf")},
               data={"holder": "Aditi Investments", "broker": "angel", "dry_run": "1"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["dry_run"] is True and d["format"] == "angel-dp" and d["scripts"] == 6 and d["count"] == 4
    assert [h["ticker"] for h in d["holdings"]] == ["ETERNAL", "TATACOMM", "TTML", "INE999Z01999"] and d["unverified"] == ["INE999Z01999"]
    assert d["nse_list"]["isins"] == 5 and d["resolved_by"]["nse-isin"] == 3
    assert c.get("/api/share-master/portfolio/csv").json()["count"] == 0                                      # dry run stored nothing
    # a .pdf with no broker field is still Angel (only Angel sends a PDF); stores 4 rows for the holder
    r = c.post("/api/share-master/portfolio/import", files={"file": ("Aditi_DP.pdf", pdf, "application/pdf")}, data={"holder": "Aditi Investments"})
    assert r.status_code == 200 and r.json()["stored"] == 4 and r.json()["format"] == "angel-dp", r.text
    csvs = c.get("/api/share-master/portfolio/csv").json()
    assert csvs["count"] == 4 and csvs["holders"]["Aditi Investments"] == {
        "uploaded_at": csvs["holders"]["Aditi Investments"]["uploaded_at"], "n": 4, "unverified": 1, "source": "angel-dp", "broker": "Angel One"}
    e = next(h for h in csvs["holdings"] if h["ticker"] == "ETERNAL")
    assert e["isin"] == "INE758T01015" and e["qty"] == 150.0 and e["cmp"] == 260.0 and e["avg_price"] is None and e["broker"] == "Angel One"
    # the same statement as .txt (pdftotext output) with ?broker=angel; raw text/plain too
    r = c.post("/api/share-master/portfolio/import?broker=angel", files={"file": ("dp.txt", ANGEL_TXT.encode(), "text/plain")}, data={"holder": "Aditi Investments"})
    assert r.status_code == 200 and r.json()["stored"] == 4
    r = c.post("/api/share-master/portfolio/import?holder=Aditi%20Investments&broker=Angel%20One", content=ANGEL_TXT.encode(), headers={"content-type": "text/plain"})
    assert r.status_code == 200 and r.json()["count"] == 4
    assert c.post("/api/share-master/portfolio/import?holder=X&broker=angel", content=b"nothing here", headers={"content-type": "text/plain"}).status_code == 400
    assert c.post("/api/share-master/portfolio/import", files={"file": ("x.pdf", b"%PDF-1.4 junk", "application/pdf")}, data={"holder": "X"}).status_code == 400
    # the CoS's Angel-derived JSON (pf/<holder>_angel_<date>.json): null ticker → ISIN column, then name
    r = c.post("/api/share-master/portfolio/import", json={"holder": "Aditi Investments", "broker": "Angel One", "as_of": "2026-10-08", "holdings": [
        {"isin": "INE758T01015", "name": "ETERNAL-EQ", "qty": 150, "value": 39000.0, "cmp": 260.0, "ticker": None, "holder": "Aditi Investments", "broker": "Angel One"},
        {"isin": "INE517B01013", "name": "TATA TELESERVICES (MAHARASHTRA)-EQ", "qty": 2000, "value": 160000.0, "cmp": 80.0, "ticker": "TTML"},
        {"isin": "INE999Z01999", "name": "SOME UNKNOWN CO-EQ", "qty": 5, "value": 1000.0, "cmp": 200.0, "ticker": None}]})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["format"] == "json" and d["stored"] == 3 and d["unverified"] == ["INE999Z01999"]
    assert {h["ticker"]: h["resolved_by"] for h in d["holdings"]} == {"ETERNAL": "nse-isin", "TTML": "given", "INE999Z01999": "unresolved"}
    assert {h["ticker"]: h["value"] for h in d["holdings"]} == {"ETERNAL": 39000.0, "TTML": 160000.0, "INE999Z01999": 1000.0}
    # Directive 18: the Angel holder shows up in the holders map, and the account status names the statement
    hold = c.get("/api/share-master/holders").json()["holders"]
    assert hold["ETERNAL"] == [{"holder": "Aditi Investments", "qty": 150.0}]
    st = c.get("/api/share-master").json()
    acc = {a["account"]: a for a in st["portfolio"]["accounts"]}
    assert acc["Aditi Investments"]["ok"] is True and acc["Aditi Investments"]["broker"] == "Angel One"
    assert acc["Aditi Investments"]["note"].startswith("Aditi Investments: Angel One DP statement uploaded ") and "1 ISIN(s) not in the NSE list" in acc["Aditi Investments"]["note"]
    # refresh: no broker session → the statement rows are the Portfolio tab, source says angel-json, cmp is the fallback price
    monkeypatch.setattr(sm, "fetch_book", lambda timeout=25.0: {"error": "connection refused"})
    monkeypatch.setattr(sm, "fetch_ltps", lambda tickers: {t: {"ETERNAL": 265.0}.get(t) for t in tickers})
    d = c.post("/api/share-master/refresh", json={}).json()
    rows = {x["ticker"]: x for x in d["portfolio"]["rows"] if x["account"] == "Aditi Investments"}
    assert rows["ETERNAL"]["ltp"] == 265.0 and rows["TTML"]["ltp"] == 80.0 and rows["TTML"]["value"] == 160000.0   # Yahoo n/a → the statement's price
    assert rows["TTML"]["source"].startswith("angel-json ") and rows["INE999Z01999"]["flag"] == "unverified ticker"
    assert rows["ETERNAL"]["invested"] is None and rows["ETERNAL"]["pnl"] is None                             # no cost in a DP statement
    assert d["ltp_missing"] == ["INE999Z01999", "TTML"]


def test_import_tool_routes_each_shape(monkeypatch, tmp_path, capsys):
    """tools/import_share_master.py: a directory of pf/*.json → the right endpoint per shape, one line
    per file with the HTTP status; --stdin takes env DOC (then stdin); --refresh alone just refreshes."""
    import importlib.util
    spec = importlib.util.spec_from_file_location("import_share_master", os.path.join(ROOT, "tools", "import_share_master.py"))
    tool = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(tool)
    calls = []

    def fake_post(path, body, timeout=240):
        calls.append((path, body))
        if path.endswith("/portfolio/import"):
            if not body.get("holdings"):
                return 400, {"detail": "no holdings parsed"}
            return 200, {"count": len(body["holdings"]), "unverified": ["INE999Z01999"] if body.get("broker") else [],
                         "resolved_by": {"nse-isin": 1}, "nse_list": {"count": 5}}
        if path.endswith("/positions/import"):
            return 200, {"stored": len(body["positions"]), "accounts": ["SA"], "rejected": [], "summary": {"pl": -1.5}}
        if path.endswith("/refresh"):
            return 200, {"vault": {"saved": True, "path": "finance/Share_Master.xlsx"}, "summary": {"total": {"n": 3}}, "ltp_missing": ["TTML"]}
        return 404, {"detail": "?"}
    monkeypatch.setattr(tool, "post", fake_post)
    (tmp_path / "aditi_2026-10-09.json").write_text(json.dumps({"holder": "Aditi", "holdings": [{"code": "ZOMATOEQ", "ticker": "ETERNAL", "qty": 1, "avg": 1, "cmp": 1, "pl_pct": float("inf")}]}))
    (tmp_path / "aditi_investments_angel_2026-10-09.json").write_text(json.dumps({"holder": "Aditi Investments", "broker": "Angel One", "holdings": [
        {"isin": "INE758T01015", "name": "ETERNAL-EQ", "qty": 150, "value": 39000.0, "cmp": 260.0, "ticker": None}]}))
    (tmp_path / "positions_2026-10-09.json").write_text(json.dumps({"positions": [{"account": "SA", "instrument": "BHEL 27-Oct-26 450 CE", "side": "BUY",
                                                                                   "qty": 1, "avg": 1, "ltp": 1, "pl": 0, "product": "OVERNIGHT"}], "as_of": "x"}))
    (tmp_path / "live_noprice.json").write_text(json.dumps({"RBLBANK": None}))
    (tmp_path / "empty.json").write_text(json.dumps({"holder": "Nobody", "holdings": []}))
    (tmp_path / "bad.json").write_text("{nope")
    rc = tool.main(["x", str(tmp_path)])
    out = capsys.readouterr().out.splitlines()
    assert rc == 1 and len(out) == 7                                              # 6 files + the refresh line
    assert out[0].startswith("aditi_2026-10-09.json: holdings Aditi → HTTP 200, 1 stored, 0 unverified")
    assert out[1].startswith("aditi_investments_angel_2026-10-09.json: holdings Aditi Investments (Angel One) → HTTP 200, 1 stored, 1 unverified (INE999Z01999)")
    assert out[2] == "bad.json: unreadable (Expecting property name enclosed in double quotes: line 1 column 2 (char 1))"
    assert out[3] == "empty.json: holdings Nobody → HTTP 400 no holdings parsed"
    assert out[4] == "live_noprice.json: skipped (neither holdings nor positions)"
    assert out[5] == "positions_2026-10-09.json: positions → HTTP 200, 1 stored for ['SA'], 0 rejected, P&L -1.5"
    assert out[6] == "refresh → HTTP 200, 3 holdings, workbook saved finance/Share_Master.xlsx, ltp n/a: TTML"
    assert [p for p, _ in calls] == ["/api/share-master/portfolio/import", "/api/share-master/portfolio/import",
                                     "/api/share-master/portfolio/import", "/api/share-master/positions/import", "/api/share-master/refresh"]
    assert calls[1][1]["broker"] == "Angel One" and calls[1][1]["holder"] == "Aditi Investments"
    assert calls[0][1]["holdings"][0]["pl_pct"] is None                          # Infinity (avg cost 0) → null, strict JSON
    assert tool.finite({"a": [float("nan"), 1.5, {"b": float("-inf")}]}) == {"a": [None, 1.5, {"b": None}]}
    # --stdin with env DOC (the VPS one-liner), no refresh; --refresh alone; --no-refresh on a directory
    calls.clear()
    monkeypatch.setenv("DOC", (tmp_path / "positions_2026-10-09.json").read_text())
    monkeypatch.setenv("DOC_NAME", "pf/positions_2026-10-09.json")
    assert tool.main(["x", "--stdin"]) == 0 and [p for p, _ in calls] == ["/api/share-master/positions/import"]
    assert capsys.readouterr().out.startswith("pf/positions_2026-10-09.json: positions → HTTP 200")
    monkeypatch.setenv("DOC", "")
    monkeypatch.setattr(sys, "stdin", io.StringIO((tmp_path / "aditi_2026-10-09.json").read_text()))
    assert tool.main(["x", "--stdin"]) == 0 and calls[-1][0] == "/api/share-master/portfolio/import"
    calls.clear()
    assert tool.main(["x", "--refresh"]) == 0 and [p for p, _ in calls] == ["/api/share-master/refresh"]
    calls.clear()
    tool.main(["x", str(tmp_path), "--no-refresh"])
    assert "/api/share-master/refresh" not in [p for p, _ in calls]
    assert tool.main(["x", str(tmp_path / "none")]) == 1
