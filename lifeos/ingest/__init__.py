"""ERP export ingest — turn a per-entity CSV/XLSX into canonical sales rows.

Each entity runs a different ERP, so headers differ. Rather than hard-code columns
we keep a per-entity *mapping* (canonical field -> that ERP's header) in
lifeos/mappings/<slug>.yaml. A file that arrives with no mapping yet is not
guessed at — the source reports "mapping needed" and lists the detected headers so
the mapping can be written in seconds (rule 1: honest, never invented).

Canonical fields (the spec handed to the ERP side):
  entity date doctype docno party item qty rate amount tax gross status owner
  due_date outstanding

A file's *report kind* (pipeline / invoiced / receivables) is inferred from its
filename so the Sales panel can total the right thing.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Optional

CANONICAL_FIELDS = [
    "entity", "date", "doctype", "docno", "party", "item", "qty", "rate",
    "amount", "tax", "gross", "status", "owner", "due_date", "outstanding",
]

# report kinds and the filename tokens that identify them
REPORT_TOKENS = {
    "pipeline": ("order", "pending", "pipeline", "quotation", "quote"),
    "invoiced": ("sales", "register", "invoice", "billed"),
    "receivables": ("outstanding", "receivable", "bills", "debtor", "ageing", "aging"),
}


@dataclass
class Table:
    headers: list[str]
    rows: list[dict[str, str]]
    path: Path
    report: str = "pipeline"


@dataclass
class CanonicalRow:
    values: dict[str, Any] = field(default_factory=dict)

    def get(self, key: str, default=None):
        return self.values.get(key, default)


def infer_report(filename: str) -> str:
    low = filename.lower()
    for kind, tokens in REPORT_TOKENS.items():
        if any(t in low for t in tokens):
            return kind
    return "pipeline"  # default: treat unknown as the funnel, the majority focus


def read_table(path: Path) -> Table:
    """Read a CSV or XLSX into a Table of string cells keyed by header. XLSX needs
    openpyxl; if it's absent for an .xlsx we raise a clear error (caught upstream)."""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        headers, rows = _read_csv(path)
    elif suffix in (".xlsx", ".xlsm"):
        headers, rows = _read_xlsx(path)
    else:
        raise ValueError(f"unsupported export type '{suffix}' (use .csv or .xlsx)")
    return Table(headers=headers, rows=rows, path=path, report=infer_report(path.name))


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        headers = [h.strip() for h in (reader.fieldnames or [])]
        rows = [{(k or "").strip(): (v or "").strip() for k, v in row.items()} for row in reader]
    return headers, rows


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError("openpyxl not installed — cannot read .xlsx exports") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    rows: list[dict[str, str]] = []
    for r in rows_iter:
        if r is None or all(c is None for c in r):
            continue
        cell = {headers[i]: ("" if v is None else str(v).strip()) for i, v in enumerate(r) if i < len(headers)}
        rows.append(cell)
    wb.close()
    return headers, rows


def apply_mapping(table: Table, mapping: dict[str, Any]) -> list[CanonicalRow]:
    """Turn raw rows into CanonicalRows using a mapping {canonical: header|{literal}}.

    A canonical field may map to a source header (string) or to a fixed value via
    {"literal": "VWLR"} — useful for `entity` when the ERP file omits it."""
    cols: dict[str, Any] = mapping.get("columns", {})
    out: list[CanonicalRow] = []
    for raw in table.rows:
        values: dict[str, Any] = {}
        for canon, spec in cols.items():
            if isinstance(spec, dict) and "literal" in spec:
                values[canon] = spec["literal"]
            elif isinstance(spec, str):
                values[canon] = raw.get(spec, "")
        out.append(CanonicalRow(values=values))
    return out


def load_mapping(slug: str, mappings_dir: Optional[Path] = None) -> Optional[dict]:
    """Load lifeos/mappings/<slug>.yaml, or None if no mapping exists yet."""
    base = mappings_dir or (Path(__file__).resolve().parent.parent / "mappings")
    path = base / f"{slug}.yaml"
    if not path.exists():
        return None
    try:
        import yaml
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("pyyaml not installed — cannot read entity mappings") from exc
    with path.open(encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def to_number(s: Any) -> float:
    """Parse an amount cell tolerant of ₹, commas, and blanks. Returns 0.0 on junk."""
    if s is None:
        return 0.0
    if isinstance(s, (int, float)):
        return float(s)
    txt = str(s).strip().replace(",", "").replace("₹", "").replace("Rs.", "").replace("Rs", "")
    if not txt:
        return 0.0
    try:
        return float(txt)
    except ValueError:
        return 0.0
